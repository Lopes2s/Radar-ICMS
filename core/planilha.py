# core/planilha.py
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Alignment

COLS_CONSULTAS = [
    'Ano', 'Nº da Consulta', 'Data da Publicação', 'Súmula',
    'Problema da Consulta', 'CNAE Detectado', 'Protocolo', 'Resposta',
]
COLS_REGIMES = [
    'ANO', 'DATA', 'Nº DOE', 'RE PR COMPETITIVO',
    'Nº DO REGIME ESPECIAL', 'EMPRESA', 'CNPJ REQUERENTE', 'CAD/ICMS',
    'CNAE REQUERENTE', 'DESCRIÇÃO CNAE', 'VIGÊNCIA DO RE', 'EMENTA',
    'ABRANGÊNCIA', 'BENEFÍCIOS/PROCEDIMENTOS', 'DISPOSIÇÕES GERAIS',
]
COLS_DESCARTES = ['Nº RE', 'Motivo', 'Empresa', 'Arquivo']
COLS_ARQUIVOS = ['Arquivo', 'Tipo', 'Consultas', 'Regimes', 'Descartes', 'Erro']

_HEADER_FILL = PatternFill('solid', fgColor='1F4E78')
_HEADER_FONT = Font(color='FFFFFF', bold=True)

# Limite do Excel por célula: 32.767 caracteres. Acima disso o arquivo abre
# com aviso de "reparo". Trunca com marca visível para o revisor saber.
_LIMITE_CELULA = 32767
MARCA_TRUNCADO = ' […texto truncado; ver PDF de origem]'


def _sanear(valor):
    """Prepara texto extraído de PDF para uma célula.

    Remove caracteres de controle (o openpyxl levanta IllegalCharacterError)
    e trunca no limite do Excel. Devolve (valor, forcar_texto): textos que
    começam com '=' viram fórmula no openpyxl, então o chamador marca a
    célula como texto.
    """
    if not isinstance(valor, str):
        return valor, False
    valor = ILLEGAL_CHARACTERS_RE.sub('', valor)
    if len(valor) > _LIMITE_CELULA:
        valor = valor[:_LIMITE_CELULA - len(MARCA_TRUNCADO)] + MARCA_TRUNCADO
    return valor, valor.startswith('=')


COLUNAS_COM_QUEBRA = {'Resposta'}
_LARGURA_COLUNA_QUEBRA = 80


def _escreve_aba(ws, cols, linhas, colunas_quebra=()):
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    for n_linha, reg in enumerate(linhas, start=2):
        for n_col, campo in enumerate(cols, start=1):
            valor, forcar_texto = _sanear(reg.get(campo, ''))
            cell = ws.cell(row=n_linha, column=n_col, value=valor)
            if forcar_texto:
                cell.data_type = 's'
            if campo in colunas_quebra:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
    for i, col in enumerate(cols, 1):
        letra = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letra].width = (
            _LARGURA_COLUNA_QUEBRA if col in colunas_quebra
            else min(max(len(col) + 2, 14), 55))


def gerar_xlsx(caminho_saida, consultas=None, regimes=None, descartes=None,
               arquivos=None):
    """Grava o .xlsx em `caminho_saida` (caminho ou objeto binário, ex. BytesIO).

    `arquivos`: uma linha por PDF processado (ver core.pipeline.Lote.arquivos).
    """
    wb = Workbook()
    wb.remove(wb.active)
    if consultas:
        _escreve_aba(wb.create_sheet('Consultas — a revisar'), COLS_CONSULTAS, consultas,
                     colunas_quebra=COLUNAS_COM_QUEBRA)
    if regimes:
        _escreve_aba(wb.create_sheet('Regimes — a revisar'), COLS_REGIMES, regimes)
    if descartes:
        _escreve_aba(wb.create_sheet('Log de descartes'), COLS_DESCARTES, descartes)
    if arquivos:
        _escreve_aba(wb.create_sheet('Arquivos processados'), COLS_ARQUIVOS, arquivos)
    if not wb.sheetnames:
        wb.create_sheet('Vazio')
    wb.save(caminho_saida)
    return caminho_saida
