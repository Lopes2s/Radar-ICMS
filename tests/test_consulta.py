# tests/test_consulta.py
import io

import pytest
from core import armazenamento, consulta


@pytest.fixture
def conn(tmp_path):
    c = armazenamento.conectar(str(tmp_path / "t.db"))
    armazenamento.criar_esquema(c)
    armazenamento.salvar_consultas(c, [
        {'Ano': '2026', 'Nº da Consulta': '001', 'Data da Publicação': '',
         'Protocolo': '', 'Súmula': 'ICMS', 'Problema da Consulta': 'madeira',
         'CNAE Detectado': ''},
        {'Ano': '2023', 'Nº da Consulta': '001', 'Data da Publicação': '',
         'Protocolo': '', 'Súmula': '', 'Problema da Consulta': '',
         'CNAE Detectado': ''},
    ])
    armazenamento.salvar_regimes(c, [
        {'ANO': '2026', 'DATA': '', 'Nº DOE': '', 'RE PR COMPETITIVO': '',
         'Nº DO REGIME ESPECIAL': '1/2026', 'EMPRESA': 'KRONA', 'CNPJ REQUERENTE': '',
         'CAD/ICMS': '', 'CNAE REQUERENTE': '', 'DESCRIÇÃO CNAE': '',
         'VIGÊNCIA DO RE': '', 'EMENTA': ''},
    ])
    return c


def test_buscar_consulta_delega_para_armazenamento(conn):
    assert len(consulta.buscar(conn, 'consulta', texto='madeira')) == 1


def test_buscar_regime_delega_para_armazenamento(conn):
    assert len(consulta.buscar(conn, 'regime', texto='KRONA')) == 1


def test_tipo_invalido_levanta_erro_de_programacao(conn):
    with pytest.raises(ValueError):
        consulta.buscar(conn, 'algo_que_nao_existe')


def test_anos_disponiveis_consulta(conn):
    assert consulta.anos_disponiveis(conn, 'consulta') == ['2023', '2026']


def test_anos_disponiveis_regime(conn):
    assert consulta.anos_disponiveis(conn, 'regime') == ['2026']


def test_anos_disponiveis_tipo_invalido_levanta_erro(conn):
    with pytest.raises(ValueError):
        consulta.anos_disponiveis(conn, 'algo_que_nao_existe')


def test_exportar_xlsx_consulta_gera_planilha_so_com_a_aba_de_consultas():
    from openpyxl import load_workbook
    resultados = [{'Ano': '2026', 'Nº da Consulta': '001', 'Data da Publicação': '',
                  'Protocolo': '', 'Súmula': 'ICMS', 'Problema da Consulta': 'x',
                  'CNAE Detectado': ''}]
    conteudo = consulta.exportar_xlsx('consulta', resultados)
    wb = load_workbook(io.BytesIO(conteudo))
    assert wb.sheetnames == ['Consultas — a revisar']
    ws = wb['Consultas — a revisar']
    assert ws.cell(row=2, column=2).value == '001'


def test_exportar_xlsx_regime_gera_planilha_so_com_a_aba_de_regimes():
    from openpyxl import load_workbook
    resultados = [{'ANO': '2026', 'DATA': '', 'Nº DOE': '', 'RE PR COMPETITIVO': '',
                  'Nº DO REGIME ESPECIAL': '1/2026', 'EMPRESA': 'KRONA',
                  'CNPJ REQUERENTE': '', 'CAD/ICMS': '', 'CNAE REQUERENTE': '',
                  'DESCRIÇÃO CNAE': '', 'VIGÊNCIA DO RE': '', 'EMENTA': ''}]
    conteudo = consulta.exportar_xlsx('regime', resultados)
    wb = load_workbook(io.BytesIO(conteudo))
    assert wb.sheetnames == ['Regimes — a revisar']


def test_exportar_xlsx_sem_resultados_ainda_gera_planilha_valida():
    from openpyxl import load_workbook
    conteudo = consulta.exportar_xlsx('consulta', [])
    load_workbook(io.BytesIO(conteudo))  # não pode levantar exceção


def test_exportar_xlsx_tipo_invalido_levanta_erro():
    with pytest.raises(ValueError):
        consulta.exportar_xlsx('algo_que_nao_existe', [])
