# tests/test_processar.py
import io
import os
import shutil
import sys

from openpyxl import load_workbook

import processar
from core import armazenamento

FIX = os.path.join(os.path.dirname(__file__), "fixtures")


def _stdout_ascii_estrito(monkeypatch):
    """Substitui sys.stdout por um stream que só aceita ASCII — o mesmo
    comportamento de um console/pipe sem suporte a UTF-8 (redirecionamento,
    Git Bash/mintty, tarefa agendada) que causava UnicodeEncodeError antes
    da correção em processar.main()."""
    buf = io.BytesIO()
    fake = io.TextIOWrapper(buf, encoding='ascii', errors='strict')
    monkeypatch.setattr(sys, 'stdout', fake)
    return fake, buf


def _aba_arquivos(saida):
    ws = load_workbook(saida)["Arquivos processados"]
    linhas = list(ws.iter_rows(values_only=True))
    cab = linhas[0]
    return [dict(zip(cab, lin)) for lin in linhas[1:]]


def test_importar_processar_nao_roda_o_bootstrap(monkeypatch):
    import importlib
    chamou = []
    monkeypatch.setattr('core.bootstrap.garantir_dependencias',
                        lambda *a, **k: chamou.append(a))
    importlib.reload(processar)
    assert chamou == []


def test_persiste_no_banco_por_padrao(tmp_path):
    pasta = tmp_path / "pdfs"
    pasta.mkdir()
    shutil.copy(os.path.join(FIX, "Consultas_1_a_3_de_2026.pdf"), pasta / "c.pdf")
    banco = tmp_path / "m.db"

    codigo = processar.main([str(pasta), "--banco", str(banco)])

    assert codigo == 0
    conn = armazenamento.conectar(str(banco))
    assert len(armazenamento.buscar_consultas(conn)) == 3


def test_xlsx_e_opcional_e_nao_obrigatorio(tmp_path):
    pasta = tmp_path / "pdfs"
    pasta.mkdir()
    shutil.copy(os.path.join(FIX, "Consultas_1_a_3_de_2026.pdf"), pasta / "c.pdf")

    codigo = processar.main([str(pasta), "--banco", str(tmp_path / "m.db")])

    assert codigo == 0
    assert not (tmp_path / "saida.xlsx").exists()  # nada gerado sem --xlsx


def test_xlsx_gerado_quando_pedido_explicitamente(tmp_path):
    pasta = tmp_path / "Consultas São José"
    pasta.mkdir()
    shutil.copy(os.path.join(FIX, "Consultas_1_a_3_de_2026.pdf"),
                pasta / "Consultas de março.pdf")
    (pasta / "corrompido.pdf").write_bytes(b"isto nao e um pdf")
    xlsx = tmp_path / "saida.xlsx"

    codigo = processar.main([str(pasta), "--banco", str(tmp_path / "m.db"),
                             "--xlsx", str(xlsx)])

    assert codigo == 0
    wb = load_workbook(xlsx)
    assert "Consultas — a revisar" in wb.sheetnames
    arquivos = {a['Arquivo']: a for a in _aba_arquivos(xlsx)}
    assert arquivos["Consultas de março.pdf"]['Tipo'] == 'consulta'
    assert arquivos["Consultas de março.pdf"]['Consultas'] == 3
    assert arquivos["corrompido.pdf"]['Tipo'] == 'erro'
    assert arquivos["corrompido.pdf"]['Erro']


def test_planilha_aberta_no_excel_da_mensagem_clara(tmp_path, monkeypatch, capsys):
    pasta = tmp_path / "pdfs"
    pasta.mkdir()
    shutil.copy(os.path.join(FIX, "Consultas_1_a_3_de_2026.pdf"), pasta / "c.pdf")

    def _bloqueado(*a, **k):
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr('core.planilha.gerar_xlsx', _bloqueado)

    codigo = processar.main([str(pasta), "--banco", str(tmp_path / "m.db"),
                             "--xlsx", str(tmp_path / "saida.xlsx")])

    assert codigo == 1
    err = capsys.readouterr().err
    assert "Não foi possível gravar" in err
    assert "aberta" in err


def test_sem_argumentos_mostra_uso(capsys):
    assert processar.main([]) == 1
    assert "Uso:" in capsys.readouterr().out


def test_pasta_sem_pdf(tmp_path, capsys):
    assert processar.main([str(tmp_path)]) == 1
    assert "Nenhum PDF" in capsys.readouterr().out


def test_flag_sem_valor_nao_quebra_com_traceback(tmp_path, capsys):
    pasta = tmp_path / "pdfs"
    pasta.mkdir()
    shutil.copy(os.path.join(FIX, "Consultas_1_a_3_de_2026.pdf"), pasta / "c.pdf")

    codigo = processar.main([str(pasta), "--banco"])

    assert codigo == 1
    assert "Uso:" in capsys.readouterr().out


def test_saida_com_acento_nao_quebra_em_console_ascii(tmp_path, monkeypatch):
    pasta = tmp_path / "pdfs"
    pasta.mkdir()
    shutil.copy(os.path.join(FIX, "Consultas_1_a_3_de_2026.pdf"), pasta / "c.pdf")
    fake, buf = _stdout_ascii_estrito(monkeypatch)

    codigo = processar.main([str(pasta), "--banco", str(tmp_path / "m.db")])

    fake.flush()
    assert codigo == 0  # sucesso não pode virar código 1 só por causa do print
    saida = buf.getvalue().decode('ascii')
    assert "Consultas" in saida  # os acentos viraram '?', mas não travou
