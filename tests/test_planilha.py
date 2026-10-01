# tests/test_planilha.py
import os
import tempfile
from openpyxl import load_workbook
from core import planilha

def test_gera_abas_e_colunas():
    consultas = [{
        'Ano': '2026', 'Nº da Consulta': '001', 'Data da Publicação': '2026-01-19',
        'Súmula': 'ICMS. DIFERIMENTO.', 'Problema da Consulta': 'x',
        'CNAE Detectado': '1623-4/00', 'Protocolo': '25.152.191-3',
    }]
    regimes = [{
        'ANO': '2026', 'DATA': '2026-08-14', 'Nº DOE': '12.196',
        'RE PR COMPETITIVO': 'Não',
        'Nº DO REGIME ESPECIAL': '9.055/2026', 'EMPRESA': 'KRONA TUBOS E CONEXÕES LTDA',
        'CNPJ REQUERENTE': '00.145.602/0001-37', 'CAD/ICMS': '09903134-40',
        'CNAE REQUERENTE': '', 'DESCRIÇÃO CNAE': '', 'VIGÊNCIA DO RE': '',
        'EMENTA': '2ª Alteração...',
    }]
    with tempfile.TemporaryDirectory() as d:
        saida = os.path.join(d, "out.xlsx")
        planilha.gerar_xlsx(saida, consultas, regimes, [])
        wb = load_workbook(saida)
        assert "Consultas — a revisar" in wb.sheetnames
        assert "Regimes — a revisar" in wb.sheetnames
        ws = wb["Consultas — a revisar"]
        # linha 1 = cabeçalho; linha 2 = primeiro registro
        assert ws.cell(row=1, column=1).value == "Ano"
        assert ws.cell(row=2, column=2).value == "001"


def _gera_e_le_consulta(problema):
    consultas = [{'Ano': '2026', 'Nº da Consulta': '001',
                  'Problema da Consulta': problema}]
    with tempfile.TemporaryDirectory() as d:
        saida = os.path.join(d, "out.xlsx")
        planilha.gerar_xlsx(saida, consultas=consultas)
        ws = load_workbook(saida)["Consultas — a revisar"]
        col = planilha.COLS_CONSULTAS.index('Problema da Consulta') + 1
        return ws.cell(row=2, column=col)


def test_caractere_de_controle_do_pdf_nao_quebra_a_planilha():
    cell = _gera_e_le_consulta("texto\x0bcom\x0ccontrole\x00")
    assert cell.value == "textocomcontrole"


def test_texto_que_comeca_com_igual_fica_como_texto_e_nao_formula():
    cell = _gera_e_le_consulta("=SOMA(1;2) citado no PDF")
    assert cell.value == "=SOMA(1;2) citado no PDF"
    assert cell.data_type == "s"


def test_texto_maior_que_o_limite_do_excel_e_truncado_com_aviso():
    cell = _gera_e_le_consulta("a" * 40000)
    assert len(cell.value) <= 32767
    assert cell.value.endswith(planilha.MARCA_TRUNCADO)


def test_coluna_resposta_presente_e_preenchida():
    consultas = [{'Ano': '2026', 'Nº da Consulta': '001',
                  'Resposta': 'Diante do exposto, o pedido é deferido.'}]
    with tempfile.TemporaryDirectory() as d:
        saida = os.path.join(d, "out.xlsx")
        planilha.gerar_xlsx(saida, consultas=consultas)
        ws = load_workbook(saida)["Consultas — a revisar"]
        col = planilha.COLS_CONSULTAS.index('Resposta') + 1
        assert ws.cell(row=1, column=col).value == 'Resposta'
        assert ws.cell(row=2, column=col).value == 'Diante do exposto, o pedido é deferido.'


def test_colunas_novas_de_regimes_presentes_e_preenchidas():
    regimes = [{'ANO': '2025', 'Nº DO REGIME ESPECIAL': '8.715/2025',
                'ABRANGÊNCIA': 'texto de abrangência',
                'BENEFÍCIOS/PROCEDIMENTOS': 'texto de procedimentos',
                'DISPOSIÇÕES GERAIS': 'texto de disposições gerais'}]
    with tempfile.TemporaryDirectory() as d:
        saida = os.path.join(d, "out.xlsx")
        planilha.gerar_xlsx(saida, regimes=regimes)
        ws = load_workbook(saida)["Regimes — a revisar"]
        for campo, esperado in [
            ('ABRANGÊNCIA', 'texto de abrangência'),
            ('BENEFÍCIOS/PROCEDIMENTOS', 'texto de procedimentos'),
            ('DISPOSIÇÕES GERAIS', 'texto de disposições gerais'),
        ]:
            col = planilha.COLS_REGIMES.index(campo) + 1
            assert ws.cell(row=1, column=col).value == campo
            assert ws.cell(row=2, column=col).value == esperado


def test_gera_planilha_em_memoria():
    import io
    buf = io.BytesIO()
    planilha.gerar_xlsx(buf, consultas=[{'Ano': '2026', 'Nº da Consulta': '001'}])
    buf.seek(0)
    assert "Consultas — a revisar" in load_workbook(buf).sheetnames


def test_aba_de_arquivos_e_descartes_com_origem():
    arquivos = [{'Arquivo': 'ruim.pdf', 'Tipo': 'erro', 'Consultas': 0,
                 'Regimes': 0, 'Descartes': 0, 'Erro': 'PdfminerException: x'}]
    descartes = [{'Nº RE': '1/2026', 'Motivo': 'revogacao', 'Empresa': 'X',
                  'Arquivo': 'doe.pdf'}]
    with tempfile.TemporaryDirectory() as d:
        saida = os.path.join(d, "out.xlsx")
        planilha.gerar_xlsx(saida, descartes=descartes, arquivos=arquivos)
        wb = load_workbook(saida)
        ws = wb["Arquivos processados"]
        assert [c.value for c in ws[1]] == planilha.COLS_ARQUIVOS
        assert ws.cell(row=2, column=2).value == 'erro'
        ws_d = wb["Log de descartes"]
        assert [c.value for c in ws_d[1]] == planilha.COLS_DESCARTES
        assert ws_d.cell(row=2, column=4).value == 'doe.pdf'
